"""
extraction.py
Turns an uploaded resume / job description file into plain text.
Supports PDF, DOCX and TXT. No knowledge of FastAPI routes or MongoDB:
raises ValueError on bad input and lets main.py turn that into an HTTP 400.
"""

import io
from pathlib import Path

from docx import Document
from fastapi import UploadFile
from pypdf import PdfReader

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt"}
MAX_FILE_SIZE_MB = 5
MIN_TEXT_LENGTH = 30  # below this, the file is effectively empty


def _extract_pdf(content: bytes) -> str:
    reader = PdfReader(io.BytesIO(content))
    if reader.is_encrypted:
        raise ValueError("PDF is password-protected. Please upload an unlocked copy.")
    return "\n".join((page.extract_text() or "") for page in reader.pages)


def _extract_docx(content: bytes) -> str:
    doc = Document(io.BytesIO(content))
    parts = [p.text for p in doc.paragraphs]
    # resumes often keep skills / experience inside tables
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                parts.append(cell.text)
    return "\n".join(parts)


def _extract_txt(content: bytes) -> str:
    try:
        return content.decode("utf-8")
    except UnicodeDecodeError:
        return content.decode("latin-1")

#takes the filename and raw bytes and returns clean text
def extract_text(filename: str, content: bytes) -> str:
    """Extract clean text from raw file bytes. Raises ValueError on any problem."""
    ext = Path(filename).suffix.lower()
    if ext not in SUPPORTED_EXTENSIONS:
        raise ValueError(
            f"Unsupported file type '{ext or 'unknown'}'. Use PDF, DOCX or TXT."
        )

    if len(content) > MAX_FILE_SIZE_MB * 1024 * 1024:
        raise ValueError(f"'{filename}' is larger than {MAX_FILE_SIZE_MB} MB.")

    try:
        if ext == ".pdf":
            text = _extract_pdf(content)
        elif ext == ".docx":
            text = _extract_docx(content)
        else:
            text = _extract_txt(content)
    except ValueError:
        raise
    except Exception as exc:  # corrupt file, bad zip, etc.
        raise ValueError(f"Could not read '{filename}': the file may be corrupted.") from exc

    # collapse blank-line noise so the LLM gets less junk
    text = "\n".join(line.strip() for line in text.splitlines() if line.strip())

    if len(text) < MIN_TEXT_LENGTH:
        raise ValueError(
            f"No readable text found in '{filename}'. "
            "If it is a scanned PDF (an image), upload a text-based version."
        )
    return text


async def extract_text_from_upload(file: UploadFile) -> str:
    """Convenience wrapper for FastAPI: reads the UploadFile, then extracts."""
    content = await file.read()
    return extract_text(file.filename or "", content)